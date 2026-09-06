import os
from pathlib import Path
from typing import List
import pandas as pd

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from src.config import CHUNK_SIZE, CHUNK_OVERLAP

def load_file(file_path: Path) -> List[Document]:
    """Loads text from PDF, DOCX, CSV, or TXT into LangChain Document format with rich metadata."""
    file_path = Path(file_path)
    ext = file_path.suffix.lower()
    docs = []

    if ext == ".pdf":
        try:
            from langchain_community.document_loaders import PyPDFLoader
            loader = PyPDFLoader(str(file_path))
            raw_docs = loader.load()
            for page_idx, doc in enumerate(raw_docs):
                doc.metadata["filename"] = file_path.name
                doc.metadata["file_type"] = "PDF"
                doc.metadata["page_number"] = doc.metadata.get("page", page_idx + 1)
                docs.append(doc)
        except Exception as e:
            print(f"[Ingestion] PyPDFLoader error: {e}. Trying fallback text extraction.")
            import pypdf
            reader = pypdf.PdfReader(file_path)
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append(Document(
                        page_content=text,
                        metadata={"filename": file_path.name, "file_type": "PDF", "page_number": i + 1}
                    ))

    elif ext == ".docx":
        try:
            import docx
            doc_file = docx.Document(file_path)
            full_text = "\n".join([p.text for p in doc_file.paragraphs if p.text.strip()])
            docs.append(Document(
                page_content=full_text,
                metadata={"filename": file_path.name, "file_type": "DOCX", "page_number": 1}
            ))
        except Exception as e:
            print(f"[Ingestion] Error loading DOCX {file_path}: {e}")

    elif ext == ".csv":
        try:
            df = pd.read_csv(file_path)
            # Convert CSV rows to structured markdown / key-value snippets
            csv_text = f"Dataset: {file_path.name}\n" + df.to_markdown(index=False)
            docs.append(Document(
                page_content=csv_text,
                metadata={"filename": file_path.name, "file_type": "CSV", "page_number": 1}
            ))
        except Exception as e:
            print(f"[Ingestion] Error loading CSV {file_path}: {e}")

    elif ext in [".txt", ".md"]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            docs.append(Document(
                page_content=content,
                metadata={"filename": file_path.name, "file_type": "TXT", "page_number": 1}
            ))
        except Exception as e:
            print(f"[Ingestion] Error reading text file {file_path}: {e}")

    return docs

def chunk_documents(documents: List[Document]) -> List[Document]:
    """Splits documents into smaller overlapping chunks and assigns chunk_id metadata."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""]
    )
    chunks = splitter.split_documents(documents)
    
    for idx, chunk in enumerate(chunks):
        fname = chunk.metadata.get("filename", "doc")
        chunk.metadata["chunk_id"] = f"{fname}_chunk_{idx}"
        
    return chunks
